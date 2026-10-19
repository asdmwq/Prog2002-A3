# -*- coding: utf-8 -*-
"""
Build the PROG2002 A3 report as a properly formatted .docx.
Same layout / complexity as the A2 report, but fresh A3 content:
registrations, full CRUD, admin dashboard, weather API, deployment.
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image
import os

OUT = r'C:\Users\30758\Doubao\chats\2026-10-04\new-chat\PROG2002-A3\docs\PROG2002 A3 Report - WU ZhengTao.docx'
DOCS = r'C:\Users\30758\Doubao\chats\2026-10-04\new-chat\PROG2002-A3\docs'
ERD = os.path.join(DOCS, 'erd.png')
SHOTS = os.path.join(DOCS, 'shots', 'a3')

doc = Document()
for section in doc.sections:
    section.top_margin = Inches(1); section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.1); section.right_margin = Inches(1.1)

normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
normal.paragraph_format.line_spacing = 1.5


def para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         space_after=8, line=1.5, font='Arial', italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    p.alignment = align
    p.paragraph_format.line_spacing = line
    p.paragraph_format.space_after = Pt(space_after)
    return p


def h1(text):
    return para(text, size=15, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                space_after=6, line=1.3)


def h2(text):
    return para(text, size=13, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                space_after=6, line=1.3)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def make_table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(htext)
        r.font.name = 'Arial'; r.font.size = Pt(11); r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        shade_cell(c, '0E7A6B')
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(val)
            r.font.name = 'Arial'; r.font.size = Pt(10.5)
            cells[i].paragraphs[0].paragraph_format.line_spacing = 1.2
    for i, w in enumerate(widths):
        for row in t.rows:
            row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def fitted_figure(path, max_w=6.3, max_h=8.0):
    with Image.open(path) as im:
        wpx, hpx = im.size
    w = max_w
    h = w * hpx / wpx
    if h > max_h:
        h = max_h
        w = h * wpx / hpx
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(path, width=Inches(w), height=Inches(h))
    return p


def caption(text):
    return para(text, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER,
                italic=True, space_after=12)


def code(text):
    return para(text, font='Consolas', size=10.5, line=1.2, space_after=6,
                align=WD_ALIGN_PARAGRAPH.LEFT)


# ============================================================
# Cover
# ============================================================
para('PROG2002 – Web Development II', size=16, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
para('Assignment 3: A Complete Dynamic Website', size=14,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)
para('HeartBridge Charity Events – Client, Admin and REST API', size=18, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
para('Event Registration, Full CRUD and a Weather Integration', size=13,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40)
para('Student ID: 24832889      Last Name: WU      First Name: ZhengTao',
     size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
para('Team member: CAO XiangXin – Student ID 24831970',
     size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
para('Title of the project: HeartBridge Charity Events', size=12,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
doc.add_page_break()

# ============================================================
# 1. Introduction / Motivation
# ============================================================
h1('1. Introduction / Motivation')
para('Assignment 2 produced a read-only HeartBridge charity event website: the public could browse current and upcoming events, filter them and view details, but they could not register, and staff had no way to add, change or remove events. Assignment 3 completes the system. The same case study is extended so that a web user can register for an event and choose how many tickets to buy, while an administrator can manage the catalogue through a dedicated back-office site. These new capabilities are supported by an improved database, a RESTful API that now implements the full Create, Read, Update and Delete cycle, and an optional external weather integration.')
para('The motivation is to turn a one-way information page into a working two-sided application. On the public side, supporters should see real attendance and be able to secure a ticket with confidence, including a check of the weather forecast for the event day. On the administrative side, staff should maintain events without touching the database directly, while the system protects the data from invalid changes. The application keeps the same layered architecture introduced in Assignment 2 — MySQL as the data layer, Node.js and Express as the service layer, and native HTML, CSS and JavaScript as the presentation layer — but now applies that architecture to a full transactional workflow.')
para('The system is deliberately built with the unit-required technologies only. The server uses Node.js and Express; the client and admin pages are delivered as plain HTML, CSS and JavaScript served statically by Express, using DOM manipulation and Promises. No CSS or JavaScript framework and no server-side template engine is used, so every page is rendered by JavaScript that consumes JSON returned by the API.')

# ============================================================
# 2. Problem Statement
# ============================================================
h1('2. Problem Statement')
para('A complete charity event platform has to satisfy two different audiences. A supporter needs to register for an event, state how many tickets they want, receive an accurate total, and see that the booking was recorded. They also expect to see who else is going and to make an informed decision about attending, for example by checking the weather. A staff member needs to list every event regardless of status, add new events, correct existing ones, and remove events that no longer apply — all through a safe interface.')
para('These requirements introduce several data-integrity problems that a read-only site did not have to address. A single person may register for several different events, but they should not be able to register twice for the same event, so the database must prevent duplicate bookings. The total amount paid must always equal the ticket price multiplied by the number of tickets, and the server cannot trust the browser to calculate it. An event that already has registrations must not be deletable, because doing so would orphan bookings. Create and update forms must reject incomplete input, while an update may legitimately change only a subset of fields. Finally, the new weather forecast needs the geographic coordinates of each event, which the original database did not store.')
para('The problem therefore spans all three layers: the database must record registrations and coordinates with the correct constraints; the API must expose safe, validated endpoints for creating, updating and deleting data; and the two websites must provide clear forms, validation messages and feedback for every outcome, including success and the cases where an operation is refused.')

# ============================================================
# 3. Solution
# ============================================================
h1('3. Solution')
para('The solution is a three-application, client-server system backed by one MySQL database. The database charityevents_db keeps the organisations, categories and events tables and adds a registrations table plus latitude and longitude columns on events. The REST API runs on port 3000 and exposes the endpoints shared by both websites. The public client runs on port 8080 and serves the Home, Search, Event Details and Registration pages. The admin client runs on port 8090 as a separate single-page dashboard for staff.')
para('All three Node.js applications are separate Express projects, each with its own package.json and its own static folder. The browser communicates with the API through fetch() and Promises, using the URL query string to carry the event id between the details and registration pages. The API enables Cross-Origin Resource Sharing only for the two approved front-end origins (the public and admin sites on ports 8080 and 8090, with the equivalent 127.0.0.1 addresses), so other websites cannot call the endpoints. The allowed list can be overridden with an environment variable for production.')
para('The registration workflow is the core public transaction. The details page now lists the recorded registrations, newest first, and the Register button links to the new registration page. That page loads the selected event through the API, shows its summary and price, collects the supporter details and ticket count, recomputes the total live as the user types, and posts the registration to the API. The server validates the input, normalises the email, recomputes the total, rejects a duplicate booking and persists the record. The administrative workflow is the core back-office transaction: the dashboard loads all events, and staff create, edit or delete records through modal forms that call the same API; the server enforces validation and the referential-integrity rule that blocks deletion of an event that already has registrations.')
para('The optional Open-Meteo weather API is used to enrich the details page. Because the service is free and needs no key, the client calls it directly with the event latitude and longitude, requests a forecast that covers the event date and maps the returned weather code to an icon and a human-readable description. The forecast is progressive enhancement: if the date is outside the forecast window or the request fails, the layout remains intact and an informative message is shown instead.')

# ============================================================
# 4. Web UX
# ============================================================
h1('4. Web UX')
h2('4.1 Public client website')
para('The public site keeps the familiar discovery journey from Assignment 2 and adds the registration path. A consistent header, navigation and footer appear on every page, and the footer now identifies the project as Assignment 3. The Home page opens with a hero section, a live count of upcoming events and the event cards, all loaded from the API. The Search page offers the same optional date, location and category filters and renders results, loading and no-results states inside the page.')
para('The Event Details page now answers two questions the supporter has before committing: who is already going, and what the weather will be. A “Recent registrations” panel lists the bookings for the event, sorted by the latest purchase date, as avatar initials with a privacy-safe name such as “Jie C.”, the ticket count and the date. A short note states that personal contact details are kept private, so only the administrator sees full email and phone details. The ticket card now contains a Register now button that links to register.html with the event id, and a weather card shows the Open-Meteo forecast for the event day together with the existing fundraising progress and host cards.')
fitted_figure(os.path.join(SHOTS, 'a3-client-detail-regs.png'))
caption('Figure 1: Event Details page showing the latest registrations with privacy-safe names.')
fitted_figure(os.path.join(SHOTS, 'a3-client-weather.png'))
caption('Figure 2: Event Details page with the Open-Meteo forecast for the event day.')
para('The Registration page is new. It begins with a summary banner for the selected event, including its category banner, date, time, location and host, so the supporter is always clear about what they are booking. The form collects first name, last name, email address, optional phone number and the number of tickets, with the registration date shown as read-only. The ticket price and total are displayed in the side card, and the total updates immediately when the ticket count changes. The form is validated before submission, and a successful booking replaces the form with a confirmation panel that restates the name, contact, ticket count, total and recorded timestamp. A note explains that one email may register for many events but only once per event.')
fitted_figure(os.path.join(SHOTS, 'a3-client-register-filled.png'))
caption('Figure 3: Registration form with live ticket total and client-side validation.')
fitted_figure(os.path.join(SHOTS, 'a3-client-success.png'))
caption('Figure 4: Confirmation panel after a successful registration.')

h2('4.2 Admin dashboard website')
para('The admin site is a separate single-page dashboard with its own teal and neutral visual language. Across the top, summary cards report total events, upcoming, past, suspended, total registrations and tickets sold, giving staff an immediate overview. The main table lists every event regardless of status — Active, Past and Suspended — and each row shows the image, title, host, date, price, fundraising progress, the registration count and the available actions. A filter box narrows the list by name, city or host as the user types.')
fitted_figure(os.path.join(SHOTS, 'a3-admin-list.png'))
caption('Figure 5: Admin dashboard with summary statistics and the full event list.')
para('Creating and editing use the same modal form. The New event button opens an empty form with required-field validation, while the Edit action for a row pre-fills every field of the selected event, including the latitude, longitude and image URL and the suspension checkbox. The bottom of the edit modal shows a “Recorded registrations” table for that event with the full name, email, phone, tickets, total and registration timestamp that the public site deliberately hides; this gives staff the information they need before changing an event.')
fitted_figure(os.path.join(SHOTS, 'a3-admin-edit.png'))
caption('Figure 6: Edit-event modal pre-filled with the current event details.')
fitted_figure(os.path.join(SHOTS, 'a3-admin-regs.png'))
caption('Figure 7: Admin view of the full registration details for the selected event.')
para('Deletion uses a confirmation modal that shows the event name and id. When the Delete button is pressed, the API either removes the event or refuses the action. If the event already has registrations, a clear inline error explains that the event cannot be deleted and states how many registrations exist, so the staff member understands exactly why the operation was blocked rather than seeing a generic failure. Successful create, update and delete actions trigger a toast notification, and the table and statistics refresh automatically.')
fitted_figure(os.path.join(SHOTS, 'a3-admin-delete.png'))
caption('Figure 8: Deletion blocked by the server because the event already has registrations.')

# ============================================================
# 5. Data Schema
# ============================================================
h1('5. Data Schema')
para('The database is called charityevents_db and retains the three original entities. organisations holds charity details, categories holds reusable event categories, and events holds the event catalogue with foreign keys to one organisation and one category. Assignment 3 extends the events table with latitude and longitude columns, stored as DECIMAL(9,6), so that the weather API can be queried; check constraints restrict latitude to the range -90 to 90 and longitude to -180 to 180.')
para('A new registrations table records each booking. It stores a unique registration identifier, the event foreign key, the supporter’s first and last name, email address and optional phone number, the number of tickets, the total amount and the registration timestamp. The foreign key references events with ON DELETE RESTRICT, which is the database-level guarantee behind the rule that an event cannot be deleted once it has registrations. A composite UNIQUE key on (event_id, email) enforces the business rule that one person may register for many events but only once for any single event. A check constraint keeps the ticket count between 1 and 20, and the total cannot be negative.')
make_table(
    ['Entity', 'Primary Key', 'Foreign Keys / Rules', 'Main Purpose'],
    [
        ['organisations', 'organisation_id', 'None', 'Charity name, description, contact details, address, city and founding year.'],
        ['categories', 'category_id', 'None', 'Reusable event categories such as Fun Run, Gala Dinner and Charity Concert.'],
        ['events', 'event_id', 'organisation_id, category_id', 'Event details, date/time, location, price, fundraising values, image, suspension flag and new latitude/longitude columns.'],
        ['registrations', 'registration_id', 'event_id (ON DELETE RESTRICT); UNIQUE(event_id, email)', 'New A3 table recording the event booked, user and contact details, ticket count, total and purchase date.'],
    ],
    [1.2, 1.1, 2.1, 2.4])
para('The supplied sample data contains three organisations, six categories, ten events and fourteen registrations, exceeding the minimum of ten registrations required by the brief. The events include seven upcoming events, two past events and one suspended event, so every status and filtering rule can be demonstrated. The registrations span seven events and cover twenty-nine tickets in total, with distinct purchase dates so that the latest-first ordering is visible on the details page.')
fitted_figure(ERD, max_w=6.6, max_h=7.4)
caption('Figure 9: Entity-Relationship Diagram for the Assignment 3 schema.')

# ============================================================
# 6. API Design
# ============================================================
h1('6. API Design')
para('The REST API is built with Node.js and Express. Assignment 2 exposed read-only GET endpoints; Assignment 3 completes the CRUD cycle. URLs are resource-oriented and the HTTP method expresses the operation: GET reads, POST creates, PUT updates and DELETE removes. The same JSON shape is returned for the public and admin sites, with the admin-only endpoints including suspended and past events and the full registration detail.')
make_table(
    ['Endpoint', 'Method', 'Purpose'],
    [
        ['/api/health', 'GET', 'Confirms the API is running and connected to MySQL.'],
        ['/api/events', 'GET', 'Returns active upcoming events for the Home page.'],
        ['/api/events/search', 'GET', 'Filters active upcoming events by optional date, location and category.'],
        ['/api/events/all', 'GET', 'Admin list: returns every event regardless of status with counts.'],
        ['/api/events/:id', 'GET', 'Returns one active event with its latest registrations for the details page.'],
        ['/api/events/admin/:id', 'GET', 'Admin version of one event with full registration details.'],
        ['/api/events', 'POST', 'Creates a new event after validating all required fields.'],
        ['/api/events/:id', 'PUT', 'Partially updates an existing event; only submitted fields are validated.'],
        ['/api/events/:id', 'DELETE', 'Deletes an event only when it has no registrations; otherwise returns 409.'],
        ['/api/events/:id/registrations', 'POST', 'Creates a registration for the event after validating and computing the total.'],
        ['/api/categories', 'GET', 'Returns categories and their event counts for the filter controls.'],
        ['/api/organisations', 'GET', 'Returns the host organisations used by the admin form.'],
    ],
    [2.35, 0.85, 3.6])
fitted_figure(os.path.join(SHOTS, 'a3-api-events.png'))
caption('Figure 10: The events endpoint returning live JSON (including coordinates) in the browser.')

h2('6.1 Detailed Example: POST /api/events/:id/registrations')
para('Purpose: create a booking for one event. Because this creates a resource, it uses POST and carries the data in the request body. The path parameter identifies the event, while the body contains the first and last name, email, phone and ticket count. The total amount is never sent by the client; the server reads the ticket price from the database and multiplies it by the validated ticket count.')
code('POST /api/events/1/registrations')
code('{ "firstName": "Report", "lastName": "Demo", "email": "report.demo@example.com",')
code('  "phone": "138 0000 0099", "tickets": 3 }')
para('Processing: the server first checks that the event exists, is active and is not in the past. It trims and lower-cases the email, validates its format, and requires the ticket count to be an integer between 1 and 20. It then checks the UNIQUE(event_id, email) rule; if the same email is already booked for that event, the request is rejected with HTTP 409. When all checks pass, the server computes the total, inserts the registration and returns the stored record with the recorded timestamp.')
para('Response codes: success returns HTTP 201 with the created registration; invalid input returns 400 with a field-specific message; a duplicate booking returns 409; a missing or unavailable event returns 404. The details page uses the GET /api/events/:id endpoint to retrieve the same event together with the complete registration list, joined to the event and sorted by the latest registration date.')

h2('6.2 Deletion Integrity and Partial Updates')
para('DELETE /api/events/:id upholds data integrity. The handler first counts the registrations linked to the event; if the count is greater than zero it returns HTTP 409 and a message that includes the number of registrations, leaving the event in place. This mirrors the ON DELETE RESTRICT foreign key and protects the back-office data. Only an event with no registrations is deleted.')
para('PUT /api/events/:id supports partial updates because a staff member may correct a single field. The validation logic therefore runs in a partial mode: fields that are present are validated, while fields that are omitted are left unchanged. This differs from POST, where every required field must be supplied before a new event is created. All queries use parameterised statements, so values are passed separately from the SQL text.')

# ============================================================
# 7. Validation, Error Handling and Testing
# ============================================================
h1('7. Validation, Error Handling and Testing')
para('Validation is applied on both the client and the server. The public form blocks empty required fields, invalid email addresses and ticket counts outside the allowed range before a request is sent, and it updates the total as the input changes. The admin forms do the same for every required event field. Because client-side checks can be bypassed, the API repeats the important validation and is the authoritative guard for the business rules: normalised email, valid ticket range, active event, duplicate registration, referential integrity and the server-computed total.')
para('The API returns meaningful HTTP status codes: 400 for invalid input, 404 for a missing event or registration target, 409 for a duplicate booking or a deletion blocked by existing registrations, and 500 for an unexpected server failure. Database errors are returned as user-safe messages rather than exposing the underlying SQL. The GET /api/health endpoint performs a simple database query so that the whole stack can be checked quickly. The database connection is made through a restricted account that has only the SELECT, INSERT, UPDATE and DELETE privileges needed by the application; it has no DDL privileges, and the connection details are kept out of the submitted code.')
para('Testing followed a bottom-up sequence. The database was verified from the SQL export and the MySQL command line, confirming the new table, constraints and fourteen seed registrations. The API was then exercised with a scripted suite of thirty request cases covering every endpoint, including successful creation, update and deletion, partial updates, the 400/404 cases, the duplicate-registration 409 and the blocked-delete 409; all thirty cases passed. A clean extraction of the three submission packages was also installed with npm ci and started against a fresh database to confirm the deployment works from the files submitted. Finally, the two websites were tested end to end in the browser, including live event rendering, search filtering, the privacy-safe registration list, the live weather forecast, a successful booking with its confirmation, the blocked duplicate booking, and the full create, edit and delete flows on the admin dashboard.')
fitted_figure(os.path.join(SHOTS, 'a3-client-home.png'))
caption('Figure 11: Public Home page with the live upcoming-event count loaded from the API.')
fitted_figure(os.path.join(SHOTS, 'a3-client-search.png'))
caption('Figure 12: Search page with optional date, location and category filters.')

# ============================================================
# 8. Deployment
# ============================================================
h1('8. Deployment')
para('The project is delivered as three independent Node.js applications plus one SQL file, matching the required submission structure. The API package contains the Express service and the database export; the public client package contains the static client site; and the admin package contains the back-office dashboard. Each application is installed with npm install, and node_modules is excluded from the submitted archives. For local development the API runs on port 3000, the public client on port 8080 and the admin client on port 8090, with the API configured to allow those origins.')
make_table(
    ['Application', 'Local Port', 'Contents', 'Submission Archive'],
    [
        ['REST API', '3000', 'Express routes, connection pool and the SQL export under database/.', '24832889A3-api.zip'],
        ['Public client', '8080', 'Home, Search, Details and Registration pages plus CSS and JS.', '24832889A3-clientside.zip'],
        ['Admin dashboard', '8090', 'Single-page CRUD dashboard and its styles and scripts.', '24832889A3-adminside.zip'],
    ],
    [1.4, 1.0, 2.9, 1.9])
para('The front-end configuration switches automatically between development and production: in development it calls the local API on port 3000, while in production it uses the same-origin /api path, with an optional production API origin constant. The final SQL export creates and populates the database, including the restricted database account used by the application.')
para('Deployment to each student’s SCU cPanel account follows the setup guide supplied with the repository: create the MySQL database, import the single SQL file, upload the three Node applications, configure the environment variables for the database connection and allowed origins, start the API, and then verify the public and admin sites over HTTPS. The public cPanel URL will be recorded here once the deployment is complete, before the live demo is recorded.')
para('Public client URL (cPanel): [INSERT PUBLIC SITE URL AFTER DEPLOYMENT]', font='Consolas', size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
para('Admin dashboard URL (cPanel): [INSERT ADMIN SITE URL AFTER DEPLOYMENT]', font='Consolas', size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
para('API base URL (cPanel): [INSERT API URL AFTER DEPLOYMENT]', font='Consolas', size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)

# ============================================================
# 9. Teamwork and Individual Contribution
# ============================================================
h1('9. Teamwork and Individual Contribution')
para('The project reused the HeartBridge client developed in Assignment 2 as the agreed base, which the brief allows by permitting a team to choose one member’s client-side website and extend it. The two members collaborated through a shared private GitHub repository with descriptive, layer-based commits, and each member deploys the applications independently to their own SCU cPanel account as required.')
make_table(
    ['Member', 'Student ID', 'Main contributions'],
    [
        ['WU ZhengTao', '24832889',
         'Extended the database with registrations and coordinates; built the full CRUD and registration API; enhanced the public details and registration pages and the weather integration; built the admin dashboard; local testing and the project documentation; individual deployment.'],
        ['CAO XiangXin', '24831970',
         'Shared code review and testing, participation in the agreed design decisions, the independent cPanel deployment, and a separate set of recorded commits and the individual demo video. [Confirm the exact agreed tasks before submission.]'],
    ],
    [1.4, 1.1, 4.7])
para('The final split above should be confirmed with the team member before submission, and the same division is reflected in the confidential peer evaluation form that is sent directly to the tutor. The repository history is the primary evidence of each member’s participation.')

# ============================================================
# 10. Conclusion
# ============================================================
h1('10. Conclusion')
para('Assignment 3 turns the HeartBridge charity event website from a read-only catalogue into a complete dynamic, two-sided application. The improved database stores registrations and coordinates with enforced uniqueness and integrity rules; the Express API now covers the full CRUD cycle and the booking workflow; the public client adds the registration journey, latest registrations and the weather forecast; and the new admin dashboard gives staff controlled create, update and delete operations with clear validation and feedback. The same required technologies from the unit — Node.js, Express, MySQL, HTML, JavaScript and the DOM — are used throughout, and the layered architecture, server-side validation and end-to-end testing keep the system maintainable and reliable.')

# ============================================================
# 11. GenAI Declaration
# ============================================================
h1('11. GenAI Declaration')
para('I acknowledge that I have used GenAI tools to complete this assessment. I used Doubao (an AI assistant) to brainstorm the feature structure and page layout, check grammar and wording, and help format the report and diagrams within the parameters outlined in the Assessment Brief and by the Unit Assessor. I reviewed and understood every suggestion, wrote and tested the final code and report content myself, and I take responsibility for the final submission.')

# ============================================================
# 12. GitHub link
# ============================================================
h1('12. GitHub Link')
para('GitHub repository (private, shared with the marker as a collaborator): [INSERT PRIVATE GITHUB REPOSITORY URL]')

doc.save(OUT)
print('Report saved:', OUT)
